# tests/test_document_parser.py
"""DOCX chunking keeps headings attached to the content they introduce."""

from docx import Document as DocxDocument

from parsers.document_parser import DocumentParser


def test_heading_is_merged_into_following_paragraph(tmp_path):
    path = tmp_path / "plan.docx"
    doc = DocxDocument()
    doc.add_paragraph("Action Items")
    doc.add_paragraph("- Alice: finalize architecture by Jan 20\n- Bob: market analysis by Feb 1")
    doc.add_heading("Metrics", level=2)
    doc.add_paragraph("Monthly recurring revenue target is $750K.")
    doc.save(path)

    chunks = [d.page_content for d in DocumentParser.parse_docx(str(path))]

    # Regression: "Action Items" used to be its own chunk, which ranked first
    # for "what are the action items" but contained no answer.
    assert not any(c.endswith("\nAction Items") for c in chunks)
    assert any("Action Items\n- Alice" in c for c in chunks)
    assert any("Metrics\nMonthly recurring revenue" in c for c in chunks)
